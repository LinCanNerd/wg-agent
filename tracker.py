"""Application log (Excel) + reply matching."""

import json
import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from sites import SITE_LABELS

COLUMNS = [
    ("Ad ID", 11),
    ("Status", 11),
    ("Site", 13),
    ("Sent at", 16),
    ("Replied at", 16),
    ("Type", 8),
    ("Title", 38),
    ("Address / district", 32),
    ("Rent €", 8),
    ("Size m²", 8),
    ("Rooms / WG", 22),
    ("Available", 22),
    ("Commute", 48),
    ("Score", 6),
    ("Contact", 16),
    ("Link", 14),
    ("Message sent", 70),
    ("Reply (preview)", 50),
    ("Conversation", 14),
]
FILL = {
    "replied": PatternFill("solid", fgColor="C6EFCE"),
    "sent": PatternFill("solid", fgColor="FFF2CC"),
    "already": PatternFill("solid", fgColor="E7E6E6"),
}
STATUS_LABEL = {"sent": "sent", "replied": "replied", "already": "contacted before"}
COL = {name: i for i, (name, _) in enumerate(COLUMNS, 1)}


def _ts(t):
    return datetime.fromtimestamp(t).strftime("%Y-%m-%d %H:%M") if t else ""


def write_excel(conn, path="applications.xlsx") -> str:
    rows = conn.execute(
        "SELECT * FROM listings WHERE status IN ('sent','replied','already') ORDER BY COALESCE(sent_at, created) DESC"
    ).fetchall()
    wb = Workbook()
    ws = wb.active
    ws.title = "Applications"
    ws.append([c for c, _ in COLUMNS])
    for i, (_, w) in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        ws.cell(1, i).font = Font(bold=True, color="FFFFFF")
        ws.cell(1, i).fill = PatternFill("solid", fgColor="305496")
    ws.freeze_panes = "C2"
    for r in rows:
        d = json.loads(r["data"])
        avail = f"{d.get('available_from') or '?'} → {d.get('available_to') or 'open-ended'}"
        rooms = d.get("flatmates") or d.get("flat_type") or ("1-Zimmer" if d.get("kind") == "studio" else "")
        ws.append(
            [
                r["id"],
                STATUS_LABEL.get(r["status"], r["status"]),
                SITE_LABELS.get(d.get("site", "wg-gesucht"), d.get("site")),
                _ts(r["sent_at"]),
                _ts(r["reply_at"]),
                "Studio" if d.get("kind") == "studio" else "WG",
                d.get("title", ""),
                d.get("address") or ", ".join(filter(None, [d.get("street"), d.get("district")])),
                d.get("rent"),
                d.get("size"),
                rooms,
                avail,
                d.get("commute_text", ""),
                r["score"],
                d.get("poster", ""),
                "open ad",
                r["draft"] or "",
                r["reply_text"] or "",
                "open chat" if r["conv_url"] else "",
            ]
        )
        row = ws.max_row
        ws.cell(row, COL["Link"]).hyperlink = d.get("url")
        ws.cell(row, COL["Link"]).font = Font(color="0563C1", underline="single")
        if r["conv_url"]:
            ws.cell(row, COL["Conversation"]).hyperlink = r["conv_url"]
            ws.cell(row, COL["Conversation"]).font = Font(color="0563C1", underline="single")
        ws.cell(row, COL["Status"]).fill = FILL.get(r["status"], PatternFill())
        for name in ("Title", "Commute", "Message sent", "Reply (preview)"):
            ws.cell(row, COL[name]).alignment = Alignment(wrap_text=True, vertical="top")
    ws.auto_filter.ref = ws.dimensions
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _norm(s):
    return re.sub(r"[^a-z0-9äöüß]", "", (s or "").lower())


def match_conversation(conv: dict, sent_rows) -> list:
    """Find which sent ad a conversation belongs to, by ad title and contact name."""
    t, n = _norm(conv.get("title")), _norm(conv.get("name"))
    hits = []
    for r in sent_rows:
        d = json.loads(r["data"])
        title, poster = _norm(d.get("title")), _norm(d.get("poster"))
        score = 0
        if t and title and (t == title or (len(t) > 12 and (t in title or title in t))):
            score += 2
        if n and poster and (poster.startswith(n) or n.startswith(poster)):  # "Anna" and "Anna Schmidt"
            score += 1
        if score:
            hits.append((score, r))
    if not hits:
        return []
    best = max(s for s, _ in hits)
    return [r for s, r in hits if s == best]


def looks_like_mine(conv_text: str, my_drafts: list[str]) -> bool:
    """The latest message shown in the inbox is the one I sent (not a reply)."""
    c = _norm(conv_text)
    for d in my_drafts:
        head = _norm(d)[:40]
        if head and head in c:
            return True
    return False
